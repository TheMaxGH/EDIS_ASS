"""
Document Engine - генерация документов в различных форматах
PDF, Markdown, Docx для отчетов и результатов работы агента
"""
from typing import Optional, Dict, Any, List
from pathlib import Path
from loguru import logger
import markdown
from datetime import datetime


class DocumentEngine:
    """
    Движок для генерации документов
    Поддерживает PDF, Markdown и Docx форматы
    """
    
    def __init__(self, output_dir: str = "outputs"):
        """
        Args:
            output_dir: Директория для сохранения документов
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)
        
        logger.info(f"DocumentEngine инициализирован: {self.output_dir}")
    
    def generate_markdown(
        self,
        content: str,
        filename: str,
        metadata: Optional[Dict[str, Any]] = None
    ) -> str:
        """
        Генерация Markdown документа
        
        Args:
            content: Содержимое документа
            filename: Имя файла (без расширения)
            metadata: Метаданные для добавления в начало
        
        Returns:
            Путь к созданному файлу
        """
        filepath = self.output_dir / f"{filename}.md"
        
        # Формируем документ
        doc_content = ""
        
        # Добавляем метаданные если есть
        if metadata:
            doc_content += "---\n"
            for key, value in metadata.items():
                doc_content += f"{key}: {value}\n"
            doc_content += "---\n\n"
        
        doc_content += content
        
        # Сохраняем
        filepath.write_text(doc_content, encoding="utf-8")
        logger.info(f"Markdown документ создан: {filepath}")
        
        return str(filepath)
    
    def generate_pdf(
        self,
        content: str,
        filename: str,
        title: Optional[str] = None,
        from_markdown: bool = True
    ) -> str:
        """
        Генерация PDF документа
        
        Args:
            content: Содержимое (Markdown или HTML)
            filename: Имя файла (без расширения)
            title: Заголовок документа
            from_markdown: Конвертировать из Markdown в HTML
        
        Returns:
            Путь к созданному файлу
        """
        try:
            from weasyprint import HTML, CSS
            from weasyprint.text.fonts import FontConfiguration
        except ImportError:
            logger.error("WeasyPrint не установлен. Установите: pip install weasyprint")
            raise ImportError("WeasyPrint требуется для генерации PDF")
        
        filepath = self.output_dir / f"{filename}.pdf"
        
        # Конвертируем Markdown в HTML если нужно
        if from_markdown:
            html_content = markdown.markdown(
                content,
                extensions=['extra', 'codehilite', 'tables', 'toc']
            )
        else:
            html_content = content
        
        # Формируем полный HTML документ
        full_html = f"""
<!DOCTYPE html>
<html>
<head>
    <meta charset="utf-8">
    <title>{title or filename}</title>
    <style>
        @page {{
            size: A4;
            margin: 2cm;
        }}
        body {{
            font-family: 'Arial', sans-serif;
            line-height: 1.6;
            color: #333;
        }}
        h1, h2, h3 {{
            color: #2c3e50;
            margin-top: 1.5em;
        }}
        h1 {{
            border-bottom: 2px solid #3498db;
            padding-bottom: 0.3em;
        }}
        code {{
            background-color: #f4f4f4;
            padding: 2px 6px;
            border-radius: 3px;
            font-family: 'Courier New', monospace;
        }}
        pre {{
            background-color: #f4f4f4;
            padding: 1em;
            border-radius: 5px;
            overflow-x: auto;
        }}
        table {{
            border-collapse: collapse;
            width: 100%;
            margin: 1em 0;
        }}
        th, td {{
            border: 1px solid #ddd;
            padding: 8px;
            text-align: left;
        }}
        th {{
            background-color: #3498db;
            color: white;
        }}
        .metadata {{
            color: #7f8c8d;
            font-size: 0.9em;
            margin-bottom: 2em;
        }}
    </style>
</head>
<body>
    <div class="metadata">
        Сгенерировано: {datetime.now().strftime("%Y-%m-%d %H:%M:%S")}
    </div>
    {html_content}
</body>
</html>
"""
        
        # Генерируем PDF
        font_config = FontConfiguration()
        html_doc = HTML(string=full_html)
        html_doc.write_pdf(filepath, font_config=font_config)
        
        logger.info(f"PDF документ создан: {filepath}")
        return str(filepath)
    
    def generate_docx(
        self,
        content: str,
        filename: str,
        title: Optional[str] = None,
        from_markdown: bool = True
    ) -> str:
        """
        Генерация DOCX документа
        
        Args:
            content: Содержимое (Markdown или plain text)
            filename: Имя файла (без расширения)
            title: Заголовок документа
            from_markdown: Парсить Markdown разметку
        
        Returns:
            Путь к созданному файлу
        """
        try:
            from docx import Document
            from docx.shared import Pt, RGBColor, Inches
            from docx.enum.text import WD_ALIGN_PARAGRAPH
        except ImportError:
            logger.error("python-docx не установлен. Установите: pip install python-docx")
            raise ImportError("python-docx требуется для генерации DOCX")
        
        filepath = self.output_dir / f"{filename}.docx"
        
        doc = Document()
        
        # Добавляем заголовок
        if title:
            heading = doc.add_heading(title, level=0)
            heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Добавляем метаданные
        metadata_para = doc.add_paragraph()
        metadata_run = metadata_para.add_run(
            f"Сгенерировано: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        )
        metadata_run.font.size = Pt(9)
        metadata_run.font.color.rgb = RGBColor(128, 128, 128)
        
        doc.add_paragraph()  # Пустая строка
        
        # Обрабатываем контент
        if from_markdown:
            # Простой парсинг Markdown
            lines = content.split('\n')
            
            for line in lines:
                line = line.strip()
                
                if not line:
                    doc.add_paragraph()
                    continue
                
                # Заголовки
                if line.startswith('# '):
                    doc.add_heading(line[2:], level=1)
                elif line.startswith('## '):
                    doc.add_heading(line[3:], level=2)
                elif line.startswith('### '):
                    doc.add_heading(line[4:], level=3)
                # Код блоки
                elif line.startswith('```'):
                    continue  # Пропускаем маркеры кода
                # Списки
                elif line.startswith('- ') or line.startswith('* '):
                    doc.add_paragraph(line[2:], style='List Bullet')
                elif line.startswith('1. ') or line.startswith('2. '):
                    doc.add_paragraph(line[3:], style='List Number')
                # Обычный текст
                else:
                    para = doc.add_paragraph(line)
                    # Обработка жирного текста **text**
                    if '**' in line:
                        para.clear()
                        parts = line.split('**')
                        for i, part in enumerate(parts):
                            run = para.add_run(part)
                            if i % 2 == 1:  # Нечетные части - жирные
                                run.bold = True
        else:
            # Просто добавляем текст как есть
            doc.add_paragraph(content)
        
        # Сохраняем
        doc.save(filepath)
        logger.info(f"DOCX документ создан: {filepath}")
        
        return str(filepath)
    
    def generate_report(
        self,
        title: str,
        sections: List[Dict[str, str]],
        format: str = "markdown",
        filename: Optional[str] = None
    ) -> str:
        """
        Генерация структурированного отчета
        
        Args:
            title: Заголовок отчета
            sections: Список секций [{title: str, content: str}, ...]
            format: Формат вывода (markdown, pdf, docx)
            filename: Имя файла (автогенерация если None)
        
        Returns:
            Путь к созданному файлу
        """
        # Генерируем имя файла если не указано
        if not filename:
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            filename = f"report_{timestamp}"
        
        # Формируем содержимое
        content = f"# {title}\n\n"
        content += f"*Дата создания: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}*\n\n"
        content += "---\n\n"
        
        for section in sections:
            section_title = section.get("title", "Без заголовка")
            section_content = section.get("content", "")
            
            content += f"## {section_title}\n\n"
            content += f"{section_content}\n\n"
        
        # Генерируем в нужном формате
        if format == "markdown":
            return self.generate_markdown(content, filename)
        elif format == "pdf":
            return self.generate_pdf(content, filename, title=title)
        elif format == "docx":
            return self.generate_docx(content, filename, title=title)
        else:
            raise ValueError(f"Неподдерживаемый формат: {format}")
    
    def list_documents(self) -> List[str]:
        """
        Список всех созданных документов
        
        Returns:
            Список путей к файлам
        """
        files = list(self.output_dir.glob("*"))
        return [str(f) for f in files if f.is_file()]
